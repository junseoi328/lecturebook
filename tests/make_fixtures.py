# 테스트용 자료 만들기: 강의 PDF(글자 2쪽 + 스캔 1쪽), PPTX, DOCX, 족보 사진
import pathlib, io
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from docx import Document

ROOT = pathlib.Path(__file__).resolve().parent
FONT = str(ROOT.parent / 'fontwork' / 'LB-Regular.ttf')
OUT = ROOT / 'fixtures'
pdfmetrics.registerFont(TTFont('KR', FONT))

def text_image(lines, w=1000, h=700):
    img = Image.new('RGB', (w, h), 'white'); d = ImageDraw.Draw(img); f = ImageFont.truetype(FONT, 34)
    for i, ln in enumerate(lines): d.text((60, 60 + i * 60), ln, fill='black', font=f)
    return img

c = canvas.Canvas(str(OUT / 'os-week5.pdf'))
pages = [
    ['5주차 프로세스 스케줄링', 'CPU 스케줄링은 준비 큐의 프로세스 중 하나를 골라 CPU를 배정한다.', 'FCFS: 먼저 온 순서대로 처리. 호위 효과(convoy effect)가 생길 수 있다.', 'SJF: 실행 시간이 가장 짧은 작업 먼저. 평균 대기 시간이 최소.'],
    ['라운드 로빈(Round Robin)', '시간 할당량(time quantum)만큼 돌아가며 실행한다.', '할당량이 너무 크면 FCFS와 같아지고, 너무 작으면 문맥 교환 비용이 커진다.', '평균 대기 시간 = 각 프로세스 대기 시간의 합 / 프로세스 수'],
]
for lines in pages:
    c.setFont('KR', 14)
    for i, ln in enumerate(lines): c.drawString(60, 780 - i * 28, ln)
    c.showPage()
img = text_image(['[스캔] 우선순위 스케줄링', '기아 상태(starvation) 발생 가능', '해결: 에이징(aging)'])
c.drawImage(ImageReader(img), 40, 300, width=500, height=350); c.showPage(); c.save()

text_image(['2025 중간고사', '1. FCFS와 SJF의 차이를 서술하시오.', '2. 라운드 로빈에서 할당량이 너무 작을 때 문제점은?', '3. 에이징이란 무엇인가?']).save(OUT / 'jokbo-2025.jpg', quality=90)

p = Presentation()
for title, body in [('6주차 동기화', '임계 구역(critical section) 문제'), ('세마포어', 'wait와 signal 연산으로 상호 배제')]:
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = title; s.placeholders[1].text = body
    s.notes_slide.notes_text_frame.text = '시험에 자주 나옴'
p.save(OUT / 'os-week6.pptx')

d = Document(); d.add_heading('7주차 교착 상태', 1); d.add_paragraph('교착 상태의 네 가지 조건: 상호 배제, 점유와 대기, 비선점, 순환 대기.'); d.save(OUT / 'os-week7.docx')
print('fixtures:', sorted(x.name for x in OUT.iterdir()))
